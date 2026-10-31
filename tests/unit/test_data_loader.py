import sqlite3
from pathlib import Path

import pytest
from openpyxl import Workbook

from datamate import BusinessDataLoader, DataLoadError, TableRef


def test_csv_discovery_and_streaming_with_bom_and_semicolon(tmp_path: Path) -> None:
    business = tmp_path / "business"
    business.mkdir()
    source = business / "sales.csv"
    source.write_text(
        "\ufeffname;quantity;note\nCà phê;2;\nTrà;1;hot\n",
        encoding="utf-8",
        newline="",
    )
    (business / "README.md").write_text("context", encoding="utf-8")

    loader = BusinessDataLoader(business)
    (table,) = loader.discover()

    assert table.columns == ("name", "quantity", "note")
    assert table.encoding == "utf-8-sig"
    assert table.delimiter == ";"
    assert list(loader.iter_rows(table)) == [
        ("Cà phê", "2", ""),
        ("Trà", "1", "hot"),
    ]


def test_csv_rejects_ragged_rows_instead_of_silently_truncating(tmp_path: Path) -> None:
    source = tmp_path / "sales.csv"
    source.write_text("item,qty\ntea,1,extra\n", encoding="utf-8")
    loader = BusinessDataLoader(tmp_path)
    (table,) = loader.discover()

    with pytest.raises(DataLoadError, match="row width mismatch"):
        list(loader.iter_rows(table))


def test_csv_falls_back_to_vietnamese_windows_encoding(tmp_path: Path) -> None:
    source = tmp_path / "sales.csv"
    source.write_bytes("name,qty\nCà phê,2\n".encode("cp1258"))
    loader = BusinessDataLoader(tmp_path)
    (table,) = loader.discover()

    assert table.encoding == "cp1258"
    assert list(loader.iter_rows(table)) == [("Cà phê", "2")]


def test_workbook_exposes_each_sheet_and_preserves_duplicate_headers(tmp_path: Path) -> None:
    business = tmp_path / "business"
    business.mkdir()
    workbook = Workbook()
    first = workbook.active
    first.title = "Sales"
    first.append(["Item", "Item"])
    first.append(["Tea", 2])
    second = workbook.create_sheet("Returns")
    second.append(["Item", "Count"])
    second.append(["Coffee", 1])
    workbook.save(business / "register.xlsx")

    loader = BusinessDataLoader(business)
    tables = loader.discover()

    assert [table.sheet_name for table in tables] == ["Sales", "Returns"]
    assert tables[0].columns == ("Item", "Item")
    assert "duplicate_header_names" in tables[0].warnings
    assert list(loader.iter_rows(tables[0])) == [("Tea", 2)]
    assert list(loader.iter_rows(tables[1])) == [("Coffee", 1)]


def test_workbook_flags_possible_title_row_without_dropping_it(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Monthly sales report"])
    sheet.append(["Item", "Quantity"])
    sheet.append(["Tea", 2])
    workbook.save(tmp_path / "report.xlsx")

    loader = BusinessDataLoader(tmp_path)
    (table,) = loader.discover()

    assert table.header_row == 1
    assert table.columns == ("Monthly sales report", "__unnamed_2")
    assert "possible_title_or_multiline_header" in table.warnings
    assert list(loader.iter_rows(table)) == [
        ("Item", "Quantity"),
        ("Tea", 2),
    ]


def test_sqlite_is_read_only_and_streams_all_tables(tmp_path: Path) -> None:
    business = tmp_path / "business"
    business.mkdir()
    database = business / "shop.sqlite"
    with sqlite3.connect(database) as connection:
        connection.execute('CREATE TABLE "order""items" (id INTEGER, amount REAL)')
        connection.executemany(
            'INSERT INTO "order""items" VALUES (?, ?)', [(1, 4.5), (2, 7.0)]
        )
        connection.execute("CREATE TABLE empty (value TEXT)")

    loader = BusinessDataLoader(business)
    tables = loader.discover()
    target = next(table for table in tables if table.sqlite_table == 'order"items')
    assert target.columns == ("id", "amount")
    assert list(loader.iter_rows(target, batch_size=1)) == [(1, 4.5), (2, 7.0)]
    with pytest.raises(sqlite3.OperationalError, match="readonly"):
        with sqlite3.connect(f"{database.as_uri()}?mode=ro", uri=True) as connection:
            connection.execute("INSERT INTO empty VALUES ('changed')")


def test_unsafe_relative_path_is_rejected(tmp_path: Path) -> None:
    loader = BusinessDataLoader(tmp_path)
    table = TableRef(
        relative_path="../outside.csv",
        source_format="csv",
        name="outside",
        columns=("value",),
        encoding="utf-8",
    )

    with pytest.raises(DataLoadError, match="Unsafe business-relative path"):
        list(loader.iter_rows(table))


def test_missing_business_directory_is_reported(tmp_path: Path) -> None:
    with pytest.raises(DataLoadError, match="does not exist"):
        BusinessDataLoader(tmp_path / "missing")
