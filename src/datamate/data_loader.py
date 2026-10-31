"""Read-only, streaming adapters for business CSV, XLSX, and SQLite files."""

from __future__ import annotations

import codecs
import csv
import sqlite3
import zipfile
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from typing import Iterator, Literal

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

SourceFormat = Literal["csv", "xlsx", "sqlite"]
SUPPORTED_SUFFIXES: dict[str, SourceFormat] = {
    ".csv": "csv",
    ".xlsx": "xlsx",
    ".sqlite": "sqlite",
    ".sqlite3": "sqlite",
    ".db": "sqlite",
}


class DataLoadError(Exception):
    """Raised when a source cannot be catalogued or read without data loss."""


@dataclass(frozen=True)
class TableRef:
    """Stable metadata for a table, sheet, or CSV file within one business."""

    relative_path: str
    source_format: SourceFormat
    name: str
    columns: tuple[str, ...]
    sheet_name: str | None = None
    header_row: int = 1
    sqlite_table: str | None = None
    encoding: str | None = None
    delimiter: str = ","
    quotechar: str = '"'
    doublequote: bool = True
    escapechar: str | None = None
    skipinitialspace: bool = False
    warnings: tuple[str, ...] = ()


class BusinessDataLoader:
    """Discover and stream tabular data without writing to source files."""

    def __init__(self, business_dir: Path | str) -> None:
        root = Path(business_dir)
        if not root.exists() or not root.is_dir():
            raise DataLoadError(f"Business directory does not exist: {root}")
        self.root = root.resolve()

    def discover(self) -> tuple[TableRef, ...]:
        """Return supported tables in deterministic path/sheet/table order."""
        tables: list[TableRef] = []
        candidates = sorted(self.root.rglob("*"), key=lambda item: item.as_posix().casefold())
        for path in candidates:
            if path.is_symlink():
                raise DataLoadError(f"Symbolic links are not allowed in business data: {path}")
            if not path.is_file():
                continue
            source_format = SUPPORTED_SUFFIXES.get(path.suffix.lower())
            if source_format is None:
                continue
            relative_path = path.relative_to(self.root).as_posix()
            if source_format == "csv":
                tables.append(self._catalog_csv(path, relative_path))
            elif source_format == "xlsx":
                tables.extend(self._catalog_workbook(path, relative_path))
            else:
                tables.extend(self._catalog_sqlite(path, relative_path))
        return tuple(tables)

    def iter_rows(
        self, table: TableRef, *, batch_size: int = 512
    ) -> Iterator[tuple[object, ...]]:
        """Stream rows as tuples aligned to ``table.columns``."""
        if batch_size < 1:
            raise ValueError("batch_size must be greater than zero")
        path = self._safe_source_path(table.relative_path)
        if table.source_format == "csv":
            return self._iter_csv_rows(path, table)
        if table.source_format == "xlsx":
            return self._iter_excel_rows(path, table)
        if table.source_format == "sqlite":
            return self._iter_sqlite_rows(path, table, batch_size)
        raise DataLoadError(f"Unsupported source format: {table.source_format}")

    def _safe_source_path(self, relative_path: str) -> Path:
        relative = PurePosixPath(relative_path)
        if (
            not relative.parts
            or relative.is_absolute()
            or any(part in {"", ".", ".."} or ":" in part for part in relative.parts)
            or "\\" in relative_path
        ):
            raise DataLoadError(f"Unsafe business-relative path: {relative_path!r}")

        candidate = self.root
        for part in relative.parts:
            candidate = candidate / part
            if candidate.is_symlink():
                raise DataLoadError(f"Symbolic links are not allowed: {relative_path!r}")
        resolved = candidate.resolve()
        if not resolved.is_relative_to(self.root) or not resolved.is_file():
            raise DataLoadError(f"Source is outside the business directory: {relative_path!r}")
        return resolved

    @staticmethod
    def _header_names(
        values: tuple[object, ...] | list[object],
    ) -> tuple[tuple[str, ...], tuple[str, ...]]:
        columns: list[str] = []
        blank_positions: list[int] = []
        for index, value in enumerate(values, start=1):
            if value is None or str(value) == "":
                columns.append(f"__unnamed_{index}")
                blank_positions.append(index)
            else:
                columns.append(str(value))
        warnings: list[str] = []
        if blank_positions:
            warnings.append("blank_header_columns:" + ",".join(map(str, blank_positions)))
        if len(set(columns)) != len(columns):
            warnings.append("duplicate_header_names")
        return tuple(columns), tuple(warnings)

    @staticmethod
    def _detect_csv_encoding(path: Path) -> str:
        for encoding in ("utf-8-sig", "cp1258"):
            decoder = codecs.getincrementaldecoder(encoding)(errors="strict")
            try:
                with path.open("rb") as stream:
                    while chunk := stream.read(1024 * 1024):
                        decoder.decode(chunk, final=False)
                    decoder.decode(b"", final=True)
                return encoding
            except UnicodeDecodeError:
                continue
            except OSError as exc:
                raise DataLoadError(f"Cannot read CSV {path}: {exc}") from exc
        raise DataLoadError(f"CSV encoding is not supported (UTF-8 or CP1258): {path}")

    @staticmethod
    def _csv_dialect(path: Path, encoding: str) -> tuple[str, str, bool, str | None, bool]:
        try:
            with path.open("r", encoding=encoding, newline="") as stream:
                sample = stream.read(8192)
        except (OSError, UnicodeDecodeError) as exc:
            raise DataLoadError(f"Cannot inspect CSV {path}: {exc}") from exc
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            return (
                dialect.delimiter,
                dialect.quotechar,
                dialect.doublequote,
                dialect.escapechar,
                dialect.skipinitialspace,
            )
        except csv.Error:
            return ",", '"', True, None, False

    def _catalog_csv(self, path: Path, relative_path: str) -> TableRef:
        encoding = self._detect_csv_encoding(path)
        delimiter, quotechar, doublequote, escapechar, skipinitialspace = self._csv_dialect(
            path, encoding
        )
        try:
            with path.open("r", encoding=encoding, newline="") as stream:
                reader = csv.reader(
                    stream,
                    delimiter=delimiter,
                    quotechar=quotechar,
                    doublequote=doublequote,
                    escapechar=escapechar,
                    skipinitialspace=skipinitialspace,
                )
                header = next((row for row in reader if any(value != "" for value in row)), None)
        except (OSError, UnicodeDecodeError, csv.Error) as exc:
            raise DataLoadError(f"Cannot read CSV header from {relative_path}: {exc}") from exc
        if header is None:
            raise DataLoadError(f"CSV has no non-empty header row: {relative_path}")
        columns, warnings = self._header_names(header)
        return TableRef(
            relative_path=relative_path,
            source_format="csv",
            name=Path(relative_path).stem,
            columns=columns,
            encoding=encoding,
            delimiter=delimiter,
            quotechar=quotechar,
            doublequote=doublequote,
            escapechar=escapechar,
            skipinitialspace=skipinitialspace,
            warnings=warnings,
        )

    def _catalog_workbook(self, path: Path, relative_path: str) -> tuple[TableRef, ...]:
        tables: list[TableRef] = []
        workbook = None
        try:
            workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
            for sheet in workbook.worksheets:
                rows = sheet.iter_rows(min_row=1, max_row=2, values_only=True)
                first_row = next(rows, None)
                second_row = next(rows, None)
                if first_row is None:
                    continue
                first_count = sum(
                    value is not None and str(value) != "" for value in first_row
                )
                second_count = sum(
                    value is not None and str(value) != "" for value in second_row or ()
                )
                if first_count == 0 and second_count >= 2:
                    header = second_row
                    header_row = 2
                    classification_warning = "blank_first_row_skipped"
                elif first_count <= 1 and second_count >= 2:
                    header = first_row
                    header_row = 1
                    classification_warning = "possible_title_or_multiline_header"
                else:
                    header = first_row
                    header_row = 1
                    classification_warning = None
                if header is None or not any(
                    value is not None and str(value) != "" for value in header
                ):
                    continue
                if header_row == 1 and second_row is not None and len(second_row) > len(header):
                    header = tuple(header) + (None,) * (len(second_row) - len(header))
                columns, warnings = self._header_names(header)
                if classification_warning:
                    warnings += (classification_warning,)
                if header_row == 1 and second_row is not None and len(second_row) > len(first_row):
                    warnings += ("row_2_exceeds_header_width",)
                tables.append(
                    TableRef(
                        relative_path=relative_path,
                        source_format="xlsx",
                        name=f"{Path(relative_path).name}:{sheet.title}",
                        columns=columns,
                        sheet_name=sheet.title,
                        header_row=header_row,
                        warnings=warnings,
                    )
                )
        except (OSError, ValueError, KeyError, zipfile.BadZipFile, InvalidFileException) as exc:
            raise DataLoadError(f"Cannot inspect workbook {relative_path}: {exc}") from exc
        finally:
            if workbook is not None:
                workbook.close()
        return tuple(tables)

    def _catalog_sqlite(self, path: Path, relative_path: str) -> tuple[TableRef, ...]:
        tables: list[TableRef] = []
        connection = self._connect_sqlite(path, relative_path)
        try:
            rows = connection.execute(
                "SELECT name FROM sqlite_master "
                "WHERE type = 'table' AND name NOT LIKE 'sqlite_%' ORDER BY name"
            ).fetchall()
            for (name,) in rows:
                info = connection.execute(f"PRAGMA table_info({self._quote_identifier(name)})")
                columns = tuple(row[1] for row in info.fetchall())
                tables.append(
                    TableRef(
                        relative_path=relative_path,
                        source_format="sqlite",
                        name=name,
                        columns=columns,
                        sqlite_table=name,
                    )
                )
        except sqlite3.Error as exc:
            raise DataLoadError(f"Cannot inspect SQLite database {relative_path}: {exc}") from exc
        finally:
            connection.close()
        return tuple(tables)

    def _iter_csv_rows(self, path: Path, table: TableRef) -> Iterator[tuple[object, ...]]:
        if table.encoding is None:
            raise DataLoadError(f"CSV encoding metadata is missing: {table.relative_path}")
        try:
            with path.open("r", encoding=table.encoding, newline="") as stream:
                reader = csv.reader(
                    stream,
                    delimiter=table.delimiter,
                    quotechar=table.quotechar,
                    doublequote=table.doublequote,
                    escapechar=table.escapechar,
                    skipinitialspace=table.skipinitialspace,
                )
                header_seen = False
                for row in reader:
                    if not header_seen:
                        if not any(value != "" for value in row):
                            continue
                        header_seen = True
                        continue
                    if not row:
                        continue
                    if len(row) != len(table.columns):
                        raise DataLoadError(
                            f"CSV row width mismatch in {table.relative_path} "
                            f"at line {reader.line_num}: expected {len(table.columns)}, got {len(row)}"
                        )
                    yield tuple(row)
        except (OSError, UnicodeDecodeError, csv.Error) as exc:
            raise DataLoadError(f"Cannot stream CSV {table.relative_path}: {exc}") from exc

    def _iter_excel_rows(self, path: Path, table: TableRef) -> Iterator[tuple[object, ...]]:
        if table.sheet_name is None:
            raise DataLoadError(f"Excel worksheet metadata is missing: {table.relative_path}")
        workbook = None
        try:
            workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
            if table.sheet_name not in workbook.sheetnames:
                raise DataLoadError(
                    f"Worksheet {table.sheet_name!r} no longer exists in {table.relative_path}"
                )
            sheet = workbook[table.sheet_name]
            for row_number, row in enumerate(
                sheet.iter_rows(min_row=table.header_row + 1, values_only=True),
                start=table.header_row + 1,
            ):
                values = tuple(row)
                if len(values) > len(table.columns) and any(
                    value is not None for value in values[len(table.columns) :]
                ):
                    raise DataLoadError(
                        f"Excel row has values beyond the header in {table.relative_path}"
                        f":{table.sheet_name} at row {row_number}"
                    )
                values = values[: len(table.columns)]
                if len(values) < len(table.columns):
                    values += (None,) * (len(table.columns) - len(values))
                if any(value is not None for value in values):
                    yield values
        except DataLoadError:
            raise
        except (OSError, ValueError, KeyError, zipfile.BadZipFile, InvalidFileException) as exc:
            raise DataLoadError(
                f"Cannot stream worksheet {table.sheet_name!r} from {table.relative_path}: {exc}"
            ) from exc
        finally:
            if workbook is not None:
                workbook.close()

    def _iter_sqlite_rows(
        self, path: Path, table: TableRef, batch_size: int
    ) -> Iterator[tuple[object, ...]]:
        if table.sqlite_table is None:
            raise DataLoadError(f"SQLite table metadata is missing: {table.relative_path}")
        connection = self._connect_sqlite(path, table.relative_path)
        cursor = None
        try:
            cursor = connection.execute(
                f"SELECT * FROM {self._quote_identifier(table.sqlite_table)}"
            )
            while rows := cursor.fetchmany(batch_size):
                for row in rows:
                    yield tuple(row)
        except sqlite3.Error as exc:
            raise DataLoadError(
                f"Cannot stream SQLite table {table.sqlite_table!r} "
                f"from {table.relative_path}: {exc}"
            ) from exc
        finally:
            if cursor is not None:
                cursor.close()
            connection.close()

    @staticmethod
    def _quote_identifier(identifier: str) -> str:
        return '"' + identifier.replace('"', '""') + '"'

    @classmethod
    def _connect_sqlite(cls, path: Path, relative_path: str) -> sqlite3.Connection:
        try:
            connection = sqlite3.connect(f"{path.as_uri()}?mode=ro", uri=True)
            connection.execute("PRAGMA query_only = ON")
            return connection
        except sqlite3.Error as exc:
            raise DataLoadError(f"Cannot open SQLite database read-only {relative_path}: {exc}") from exc
